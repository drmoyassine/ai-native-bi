"""Connect data C1 (docs/connect-data-spec.md): file sources through the whole flow, on made-up exports (DEC-011). A file whose columns
have the dataset's own names loads at once; any other layout waits in the draft schema, profiled, until a person approves its mapping;
layouts are kept apart, Arabic headers stay two columns, and a file that fails the tests does not count."""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

import openpyxl
import psycopg
import pytest
from psycopg_pool import AsyncConnectionPool

from bi import accounts
from bi.connect import files, store
from bi.connect.files import Storage, load_now, process_sync
from bi.db import migrate

from .conftest import ADMIN, PoolFactory, database_url, fresh_database
from .fakes import FALL, HEADER, SPRING, csv_file, students, xlsx

AS_OF = dt.date(2026, 10, 20)
DB = 'bi_connect_test'
# An export from another system: a title above the header, other names, an email column and an Arabic header.
EXPORT = ['Student Number', 'Semester', 'Semester Label', 'Start Date', 'Major', 'Student Email', 'المستوى']


def export(path: Path, rows: list[list[Any]]) -> Path:
    book = openpyxl.Workbook()
    ws = book.active
    assert ws is not None
    ws.append(['Example Polytechnic — Students'])
    ws.append([])
    ws.append(EXPORT)
    for row in rows:
        ws.append(row)
    book.save(path)
    return path


def export_rows(n: int = 6, term: str = 'F26') -> list[list[Any]]:
    return [[f'S-{i}', term, 'Fall 2026', dt.date(2026, 8, 24), 'BSME' if i % 2 else 'BSCE', f'student{i}@example.edu', 'بكالوريوس']
            for i in range(1, n + 1)]  # fmt: skip


def same_file_read_only(kept: Path, original: Path) -> bool:
    return kept.read_bytes() == original.read_bytes() and kept.stat().st_mode & 0o777 == 0o400


def query(sql: str) -> list[tuple[Any, ...]]:
    with psycopg.connect(database_url(DB)) as conn:
        return conn.execute(sql).fetchall()


@pytest.fixture(scope='module')
async def pool(make_pool: PoolFactory) -> AsyncConnectionPool:
    p = await make_pool(DB)
    await accounts.create_user(p, 'admin@example.edu', 'Admin', 'a-long-password-0123', is_admin=True)
    return p


@pytest.fixture
def storage(tmp_path: Path) -> Storage:
    return Storage(f'file://{tmp_path}/files')


async def admin_id(pool: AsyncConnectionPool) -> str:
    async with pool.connection() as conn:
        row = await (await conn.execute("select id::text from app_user where email = 'admin@example.edu'")).fetchone()
    assert row is not None
    return str(row[0])


async def test_a_file_with_the_datasets_own_names_loads_at_once_into_raw_as_written(
    pool: AsyncConnectionPool, storage: Storage, tmp_path: Path
) -> None:
    path = xlsx(tmp_path / 'fall.xlsx', students(FALL, (1001, 'BSME', 'Bachelor'), (1002, 'BSCE', 'Bachelor')))
    ended = await load_now(pool, storage, database_url(DB), path, source_name='Registrar exports', dataset='enrollment', as_of=AS_OF)
    assert ended['status'] == 'loaded' and ended['rows'] == 2 and ended['tests'] > 0 and ended['found_header_row'] == 1
    columns = [r[0] for r in query("select column_name from information_schema.columns where table_schema = 'raw' and "
                                    "table_name = 'registrar_exports' order by ordinal_position")]  # fmt: skip
    assert 'Student ID' in columns and 'Term Code' in columns  # as written (DEC-031)
    assert query('select student_id, program, level from analytics.enrollment order by 1') == [
        ('1001', 'BSME', 'bachelor'),
        ('1002', 'BSCE', 'bachelor'),
    ]
    again = await load_now(pool, storage, database_url(DB), path, source_name='Registrar exports', dataset='enrollment', as_of=AS_OF)
    assert again['status'] == 'loaded' and again['repeat'] is True
    assert query('select count(*) from raw.registrar_exports') == [(2,)]


async def test_another_layout_waits_profiled_until_a_person_approves_its_mapping(
    pool: AsyncConnectionPool, storage: Storage, tmp_path: Path
) -> None:
    path = export(tmp_path / 'banner.xlsx', export_rows())
    ended = await load_now(pool, storage, database_url(DB), path, source_name='Banner export', dataset='enrollment', as_of=AS_OF)
    assert ended['status'] == 'review' and 'Check which column is which' in ended['note']
    assert query("select to_regclass('raw.banner_export') is null") == [(True,)]  # nothing in raw before approval
    assert query('select count(*) from draft.banner_export') == [(6,)]

    seen = await files.review(pool, ended['id'])
    assert seen is not None and seen['header']['row'] == 3 and seen['header']['above_header'] == ['Example Polytechnic — Students']
    columns = {c['header']: c for c in seen['columns']}
    assert list(columns) == EXPORT  # as written, Arabic included
    assert columns['Student Email']['kind'] == 'personal' and columns['Student Email']['summary'].startswith('personal data like a')
    assert 'student1' not in str(seen)  # no value of a personal column, anywhere
    assert columns['Start Date']['kind'] == 'date' and '2026-08-24' in columns['Start Date']['summary']
    assert seen['mapping_status'] == 'proposed'
    assert seen['mapping']['student_id'] == 'Student Number' and seen['mapping']['term_code'] == 'Semester'
    assert seen['mapping']['term_start'] == 'Start Date' and seen['mapping']['level'] == 'المستوى'
    assert {p['label'] for p in seen['parts']} >= {'Student ID', 'Term code', 'Term start date'}

    me = await admin_id(pool)
    with pytest.raises(ValueError, match='no column is chosen for: Term code'):
        await files.approve(pool, ended['id'], {**seen['mapping'], 'term_code': None}, me, 'admin@example.edu')
    await files.approve(pool, ended['id'], seen['mapping'], me, 'admin@example.edu')
    assert await process_sync(pool, storage, database_url(DB), ended['id'], 5) == 'loaded'
    assert query("select to_regclass('draft.banner_export') is null") == [(True,)]
    raw_columns = [r[0] for r in query("select column_name from information_schema.columns where table_schema = 'raw' and "
                                        "table_name = 'banner_export'")]  # fmt: skip
    assert 'المستوى' in raw_columns and 'Student Email' in raw_columns
    by_term = dict(query('select term_code, count(*) from analytics.enrollment group by 1'))
    assert by_term['F26'] == 6 and by_term['202610'] == 2  # two sources, two layouts, one dataset
    assert query("select level from analytics.enrollment where student_id = 'S-1'") == [('بكالوريوس',)]

    # The same source, a changed export (a column renamed): it waits again, starting from what was approved.
    renamed = export(tmp_path / 'banner2.xlsx', export_rows(term='S27'))
    ws_book = openpyxl.load_workbook(renamed)
    ws = ws_book.active
    assert ws is not None
    ws.cell(row=3, column=5, value='Programme')
    ws_book.save(renamed)
    changed = await load_now(pool, storage, database_url(DB), renamed, source_name='Banner export', dataset='enrollment', as_of=AS_OF)
    assert changed['status'] == 'review'
    again = await files.review(pool, changed['id'])
    assert again is not None and again['mapping']['student_id'] == 'Student Number' and again['mapping']['program'] == 'Programme'


async def test_two_arabic_headers_stay_two_columns(pool: AsyncConnectionPool, storage: Storage, tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    ws = book.active
    assert ws is not None
    ws.append(['اسم البرنامج', 'رقم البرنامج'])
    for i in range(3):
        ws.append([f'Program {i}', f'P{i}'])
    book.save(tmp_path / 'ar.xlsx')
    ended = await load_now(pool, storage, database_url(DB), tmp_path / 'ar.xlsx', source_name='Programs', dataset=None, as_of=AS_OF)
    assert ended['status'] == 'loaded'  # a raw-only source loaded by an operator is approved when made
    assert query('select "اسم البرنامج", "رقم البرنامج" from raw.programs order by 2') == [('Program 0', 'P0'), ('Program 1', 'P1'),
                                                                                            ('Program 2', 'P2')]  # fmt: skip


async def test_a_file_that_fails_the_tests_does_not_count_and_says_why(pool: AsyncConnectionPool, storage: Storage, tmp_path: Path) -> None:
    before = query('select count(*) from analytics.enrollment')
    dup = xlsx(tmp_path / 'dup.xlsx', students(SPRING, (2001, 'BSME', 'Bachelor'), (2001, 'BSCE', 'Bachelor')))
    ended = await load_now(pool, storage, database_url(DB), dup, source_name='Registrar exports', dataset='enrollment', as_of=AS_OF)
    assert ended['status'] == 'failed' and 'unique' in ended['error']
    assert query('select count(*) from analytics.enrollment') == before  # rebuilt without it


async def test_what_cannot_be_read_fails_with_a_reason(pool: AsyncConnectionPool, storage: Storage, tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    ws = book.active
    assert ws is not None
    ws.append(['Just a title'])
    book.save(tmp_path / 'title.xlsx')
    ended = await load_now(pool, storage, database_url(DB), tmp_path / 'title.xlsx', source_name='Odd', dataset=None, as_of=AS_OF)
    assert ended['status'] == 'failed' and 'header' in ended['error']
    with pytest.raises(store.SourceError, match='already a source'):
        await store.create_source(pool, name='odd', type='file', dataset=None, owner=None)
    with pytest.raises(store.SourceError, match='not available yet'):
        await store.create_source(pool, name='Banner DB', type='oracle', dataset=None, owner=None)


async def test_uploads_become_syncs_of_a_file_source_to_load_again() -> None:
    url = await fresh_database('bi_connect_migration_test')
    pool = AsyncConnectionPool(url, open=False)
    await pool.open()
    try:
        async with pool.connection() as conn:
            await conn.execute('create table schema_migration (name text primary key, applied_at timestamptz not null default now())')
            await conn.execute("insert into schema_migration (name) values ('0010_connect.sql')")
        await migrate(pool)  # everything before Connect data
        await accounts.create_user(pool, 'ie@example.edu', 'IE', 'a-long-password-0123')
        async with pool.connection() as conn:
            await conn.execute(
                """insert into upload (user_id, dataset, as_of, file_name, incoming_url, status, rows, error)
                   select id, 'enrollment', '2026-10-20', f, 'file:///x/' || f, s, 10, e from app_user,
                   (values ('a.xlsx', 'loaded', null), ('b.xlsb', 'failed', 'dbt build failed')) v(f, s, e)"""
            )
            await conn.execute("delete from schema_migration where name = '0010_connect.sql'")
        assert await migrate(pool) == ['0010_connect.sql']
        async with pool.connection() as conn:
            sources = await (await conn.execute('select name, slug, state from source')).fetchall()
            syncs = await (await conn.execute('select file_name, status, note, error from sync order by file_name')).fetchall()
            stream = await (await conn.execute('select dataset, raw_table from source_stream')).fetchall()
            gone = await (await conn.execute("select to_regclass('upload') is null, to_regclass('column_mapping') is null")).fetchone()
        assert sources == [('Enrollment files', 'enrollment_files', 'live')] and stream == [('enrollment', 'enrollment_files')]
        assert syncs[0][:2] == ('a.xlsx', 'stale') and 'load it again' in syncs[0][2]
        assert syncs[1][1:] == ('failed', None, 'dbt build failed') and gone == (True, True)
    finally:
        await pool.close()
    assert ADMIN  # the migration test ran against the test server


async def test_the_file_is_kept_unchanged_every_column_is_kept_and_a_corrected_term_replaces_it(
    pool: AsyncConnectionPool, storage: Storage, tmp_path: Path
) -> None:
    rows = students(FALL, (3001, 'BSME', 'Bachelor'), (3002, 'BSCE', 'Bachelor'))
    path = xlsx(tmp_path / 'fall.xlsx', [[*r, None] for r in rows], header=[*HEADER, 'Advisor'])
    ended = await load_now(pool, storage, database_url(DB), path, source_name='Advising', dataset='enrollment', as_of=AS_OF)
    assert ended['status'] == 'loaded'
    kept = Path(f'{storage.files_url.removeprefix("file://")}/advising/{ended["sha256"]}.xlsx')
    assert same_file_read_only(kept, path)
    assert query('select "Student ID", "Advisor", _bi_row, _bi_as_of, _bi_file_name from raw.advising order by _bi_row') == [
        ('3001', None, 2, AS_OF, 'fall.xlsx'),
        ('3002', None, 3, AS_OF, 'fall.xlsx'),
    ]  # a column empty in this file is still a column

    corrected = students(FALL, (3001, 'BSME', 'Bachelor'), (3002, 'BSCE', 'Bachelor'), (3003, 'MSDS', 'Master'))
    later = dt.date(2026, 10, 25)
    path2 = csv_file(tmp_path / 'fall-corrected.csv', [[*r, None] for r in corrected], header=[*HEADER, 'Advisor'])
    assert (await load_now(pool, storage, database_url(DB), path2, source_name='Advising', dataset='enrollment', as_of=later))[
        'status'
    ] == 'loaded'
    assert query("select count(*), max(as_of) from analytics.enrollment where term_code = '202610'") == [(3, later)]  # the term is replaced
